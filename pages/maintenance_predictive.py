# -*- coding: utf-8 -*-
"""
Onglet « Maintenance prédictive » — pannes probables de la semaine S,
plan de contrôle (top 50), résultats réels de S-1 et évolution des
5 modèles. Rien n'est calculé ni affiché avant le clic sur
« Lancer la prédictive ».

Calendrier :
  · plan de la semaine S : disponible dès le lundi de S, calculé avec les
    données antérieures à ce lundi, puis figé ;
  · résultats réels de S : affichés à partir du lundi S+1, dès que
    l'extraction couvre toute la semaine S.
"""
import io
import pandas as pd
import streamlit as st

from core import predictive_pannes as pp

JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]


def _fmt(ts):
    ts = pd.Timestamp(ts)
    return f"{JOURS[ts.weekday()]} {ts.strftime('%d/%m/%Y')}"


def _excel_plan(df, titre):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name="Plan de contrôle", startrow=2)
        ws = xw.sheets["Plan de contrôle"]
        ws["A1"] = titre
        from openpyxl.styles import Font, PatternFill, Alignment
        ws["A1"].font = Font(bold=True, size=13)
        for c in ws[3]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1E3A5F")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        largeurs = [6, 30, 36, 14, 12, 11, 10, 10, 13, 24, 52, 40]
        for i, w in enumerate(largeurs):
            ws.column_dimensions[chr(65 + i)].width = w
    buf.seek(0)
    return buf.getvalue()


def _cycle(df_full, avis_complet, fichier_date, perimetre="general"):
    """Exécute le cycle une fois par (périmètre, extraction, jour) et le garde en session."""
    aujourd_hui = pd.Timestamp.today().normalize()
    try:
        date_ext = pd.to_datetime(fichier_date, format="%d/%m/%Y")
    except Exception:
        date_ext = pd.to_datetime(fichier_date, dayfirst=True, errors="coerce")
    cle = (perimetre, str(fichier_date), str(aujourd_hui.date()))
    caches = st.session_state.setdefault("_pred_cache_perim", {})
    cache = caches.get(perimetre)
    if cache and cache["cle"] == cle:
        return cache
    reg, source = pp.charger_registre(perimetre)
    res, modifie = pp.executer_cycle(df_full, avis_complet, aujourd_hui, date_ext, reg, perimetre=perimetre)
    msg = pp.sauver_registre(reg, perimetre) if modifie else ""
    cache = {"cle": cle, "res": res, "reg": reg, "source": source, "msg": msg,
             "aujourd_hui": aujourd_hui, "date_ext": date_ext}
    caches[perimetre] = cache
    return cache


def _journal_df(journal):
    lignes = []
    for cle in sorted(journal):
        j = journal[cle]
        pct = lambda v: (round(v * 100, 1) if v is not None else None)
        lignes.append({
            "Semaine S": cle,
            "Enregistré le": str(j.get("enregistre_le", ""))[:16],
            "Semaine évaluée": j.get("semaine_evaluee", ""),
            "Source": ("Réel" if j.get("source") == "réel" else "Rétro-test")
                      + ("" if j.get("statut") == "complet" else " (provisoire)"),
            "Efficacité — précision (%)": pct(j.get("precision")),
            "Pannes détectées (%)": pct(j.get("detection")),
            "Moy. 4 sem. précision (%)": pct(j.get("precision_moy4")),
            "Δ vs semaine préc. (pts)": pct(j.get("delta_vs_prec")),
            "Δ vs 1er enregistrement (pts)": pct(j.get("delta_vs_debut")),
            "Détectées / pannes réelles": (f"{j.get('detectees')} / {j.get('pannes_reelles')}"
                                           if j.get("pannes_reelles") is not None else ""),
        })
    return pd.DataFrame(lignes)


def _excel_simple(df, feuille, titre):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name=feuille, startrow=2)
        ws = xw.sheets[feuille]
        ws["A1"] = titre
        from openpyxl.styles import Font, PatternFill, Alignment
        ws["A1"].font = Font(bold=True, size=13)
        for c in ws[3]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1E3A5F")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i in range(len(df.columns)):
            ws.column_dimensions[chr(65 + i)].width = 18
    buf.seek(0)
    return buf.getvalue()


def _section_efficacite(reg, evol, cle_S, lundi_S):
    st.markdown("---")
    st.markdown(f"#### 📊 Efficacité de la prédictive — enregistrée au début de chaque semaine")
    journal = reg.get("journal_efficacite", {})
    if not evol:
        st.info("Pas encore de semaine évaluée : l'efficacité sera enregistrée dès la première évaluation.")
        return

    crit = st.radio("Mesure de l'efficacité", ["Précision du plan", "Pannes détectées"],
                    horizontal=True, key="pred_crit_eff")
    k = "precision" if crit == "Précision du plan" else "detection"

    ent = journal.get(cle_S)
    cles = sorted(journal)
    ant = journal[cles[cles.index(cle_S) - 1]] if ent is not None and cles.index(cle_S) > 0 else None
    vals = [e["modeles"]["M5"][k] * 100 for e in evol]
    td = pp.tendance(vals)

    e1, e2, e3, e4 = st.columns(4)
    if ent is not None:
        v = ent[k] * 100
        d = (f"{(ent[k] - ant[k]) * 100:+.1f} pts vs {ant['semaine'].split('-')[1]}"
             if ant is not None else "1er enregistrement")
        e1.metric(f"Efficacité enregistrée — début {cle_S.split('-')[1]}", f"{v:.0f} %", d,
                  delta_color="normal" if ant is not None else "off")
        e1.caption(f"Plan de {ent['semaine_evaluee']} · "
                   f"{'réel' if ent.get('source') == 'réel' else 'rétro-test'}"
                   f"{'' if ent.get('statut') == 'complet' else ' · provisoire'}")
    m4 = sum(vals[-4:]) / len(vals[-4:])
    e2.metric("Moyenne des 4 dernières semaines", f"{m4:.0f} %")
    if td:
        e3.metric("Tendance", f"{td['pente']:+.1f} pts / semaine", td["sens"],
                  delta_color="normal" if td["sens"] != "stable" else "off")
        rel = f"{td['gain_rel']:+.0f} % relatif" if td["gain_rel"] is not None else ""
        e4.metric(f"Amélioration ({td['n_bloc']} dern. sem. vs {td['n_bloc']} prem.)",
                  f"{td['gain_pts']:+.1f} pts", rel)

    try:
        import plotly.graph_objects as go
        x = [e["semaine"].split("-")[1] for e in evol]
        reel = [e.get("type") == "réel" for e in evol]
        fig = go.Figure()
        fig.add_trace(go.Bar(
            x=x, y=vals, name="Efficacité hebdomadaire (M5)",
            marker_color=["#2a78d6" if r else "#b8cfee" for r in reel],
            text=[f"{v:.0f}%" for v in vals], textposition="outside"))
        mob = [sum(vals[max(0, i - 3):i + 1]) / len(vals[max(0, i - 3):i + 1]) for i in range(len(vals))]
        fig.add_trace(go.Scatter(x=x, y=mob, name="Moyenne glissante 4 sem.", mode="lines",
                                 line=dict(color="#0b0b0b", width=2)))
        if td:
            fig.add_trace(go.Scatter(
                x=x, y=td["droite"], name=f"Tendance ({td['pente']:+.1f} pts/sem.)", mode="lines",
                line=dict(color="#0ca30c" if td["pente"] >= 0 else "#d03b3b", width=2, dash="dash")))
        fig.update_layout(height=360, yaxis_title=f"{crit} (%)", plot_bgcolor="white",
                          yaxis=dict(range=[0, max(100, max(vals) + 10)], gridcolor="#eceae4"),
                          margin=dict(l=10, r=10, t=30, b=10), legend=dict(orientation="h", y=-0.18),
                          bargap=0.35)
        st.plotly_chart(fig, use_container_width=True)
    except Exception:
        pass
    st.caption("Barres claires : semaines rejouées (rétro-test) · barres foncées : plans réellement publiés "
               "puis comparés aux pannes réelles. Précision = part des postes techniques du plan tombés en panne ; "
               "pannes détectées = part des pannes réelles qui étaient dans le plan. "
               "La tendance est la pente de la régression linéaire sur toutes les semaines évaluées.")

    if journal:
        st.markdown("**Journal d'efficacité** — une ligne figée par semaine S, au lundi de publication du plan")
        jdf = _journal_df(journal)
        st.dataframe(jdf, hide_index=True, use_container_width=True)
        st.download_button(
            "📥 Télécharger le journal d'efficacité (Excel)",
            data=_excel_simple(jdf, "Journal efficacité", "Journal d'efficacité de la maintenance prédictive"),
            file_name="Journal_efficacite_predictive.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="dl_journal_eff")


def render_maintenance_predictive_tab(df_full, avis_complet, vp, fichier_date):
    st.markdown("### 🔮 Maintenance prédictive — pannes probables de la semaine")

    aujourd_hui = pd.Timestamp.today().normalize()
    lundi_S = pp.lundi(aujourd_hui)
    lundi_S1 = lundi_S + pd.Timedelta(days=7)
    s_lab = pp.semaine_label(lundi_S)
    c1, c2, c3 = st.columns(3)
    c1.metric("Semaine S", s_lab.split("-")[1], f"{lundi_S.strftime('%d/%m')} → {(lundi_S + pd.Timedelta(days=6)).strftime('%d/%m')}", delta_color="off")
    c2.metric("Plan de contrôle publié", lundi_S.strftime("%d/%m/%Y"), "1er jour de la semaine S", delta_color="off")
    c3.metric("Résultats réels de S", lundi_S1.strftime("%d/%m/%Y"), "1er jour de la semaine S+1", delta_color="off")

    if not st.session_state.get("predictive_lancee"):
        st.info("ℹ️ Les calculs, le plan de contrôle et les résultats restent masqués "
                "jusqu'au lancement de la prédictive.")
        if st.button("🔮 Lancer la prédictive", type="primary", use_container_width=True,
                     key="btn_lancer_predictive"):
            st.session_state["predictive_lancee"] = True
            st.rerun()
        return

    if df_full is None or df_full.empty or avis_complet is None or avis_complet.empty:
        st.warning("Données OT / avis indisponibles.")
        return

    perims = [k for k, v in pp.PERIMETRES.items() if not v.get("masque")]
    perimetre = st.radio("Plan de prédictive", perims, horizontal=True, key="pred_perimetre",
                         format_func=lambda k: pp.PERIMETRES[k]["label"])
    pref = pp.PERIMETRES[perimetre]["prefixes"]
    if pref:
        st.caption(f"Pannes, historique et plan limités aux postes de travail {', '.join(pref)}… "
                   "Chaque plan a ses propres modèles, son plan de contrôle et son journal d'efficacité.")

    with st.spinner("Entraînement des 5 modèles et calcul du plan de la semaine… (≈ 1 min la 1re fois par extraction)"):
        try:
            cache = _cycle(df_full, avis_complet, fichier_date, perimetre)
        except Exception as e:
            st.error(f"❌ Échec du calcul prédictif : {e}")
            return
    res, reg = cache["res"], cache["reg"]
    D, W_S, cle_S = res["D"], res["W_S"], res["cle_S"]
    pred = reg["predictions"][cle_S]
    date_ext = cache["date_ext"]

    if pd.notna(date_ext) and date_ext < lundi_S - pd.Timedelta(days=3):
        st.warning(f"⚠️ L'extraction s'arrête au {date_ext.strftime('%d/%m/%Y')} : le plan de {cle_S} "
                   "s'appuie sur des données incomplètes de la semaine précédente.")

    # ── Périmètre (postes de travail sélectionnés) ──
    postes = [str(p) for p in (vp or [])]
    if pref:   # garder les postes sélectionnés qui font partie du plan, sinon tous
        postes = [p for p in postes if p.strip().upper().startswith(pref)]
    info = D.info
    m5 = pd.Series(pred["scores"]["M5"])
    dans = [e for e in m5.index if str(info["poste_travail"].get(e, "")) in set(postes)] if postes else list(m5.index)
    m5_p = m5.reindex(dans).dropna()

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Pannes probables cette semaine", f"{m5_p.sum():.0f}", "postes techniques (périmètre)", delta_color="off")
    k2.metric("Postes techniques surveillés", f"{len(m5_p)}", f"{pred['univers']} au total", delta_color="off")
    evol = res["evolution"]
    rec = evol[-8:]
    if rec:
        prec = sum(e["modeles"]["M5"]["precision"] for e in rec) / len(rec) * 100
        det = sum(e["modeles"]["M5"]["detection"] for e in rec) / len(rec) * 100
        k3.metric("Précision du plan (M5)", f"{prec:.0f} %", f"{len(rec)} dernières semaines", delta_color="off")
        k4.metric("Pannes détectées (M5)", f"{det:.0f} %", "des pannes réelles", delta_color="off")

    # ══ 1. Plan de contrôle ══════════════════════════════════════════════
    st.markdown("---")
    suffixe = (" · " + pp.PERIMETRES[perimetre]["court"]) if pp.PERIMETRES[perimetre].get("court") else ""
    st.markdown(f"#### 🛠️ Plan de contrôle{suffixe} — {cle_S} · publié le {_fmt(lundi_S)}")
    st.caption(f"Les {pp.TOP_N} postes techniques du périmètre au risque de panne le plus élevé. "
               f"Calculé avec les données jusqu'au {pd.Timestamp(pred['debut']).strftime('%d/%m/%Y')} (exclu), "
               f"puis figé : il ne change plus pendant la semaine.")
    fams = list(pp.CONTROLES.keys())
    sel_fam = st.multiselect("Familles de défaillance", fams, default=fams, key="pred_familles",
                             help="Mécanique, électrique et instrumentation — les fuites sont exclues. "
                                  "Filtrer sur une famille donne son propre top 50.")
    plan = pp.plan_controle(D, pred["scores"]["M5"], W_S, postes=postes or None,
                            familles=(sel_fam if set(sel_fam) != set(fams) else None))
    if plan.empty:
        st.info("Aucun poste technique à risque dans le périmètre sélectionné.")
    else:
        st.dataframe(
            plan, use_container_width=True, hide_index=True, height=520,
            column_config={
                "Probabilité de panne": st.column_config.ProgressColumn(
                    "Probabilité de panne", format="%.0f %%", min_value=0, max_value=100),
                "Contrôle recommandé": st.column_config.TextColumn(width="large"),
            })
        st.download_button(
            "📥 Télécharger le plan de contrôle (Excel)",
            data=_excel_plan(plan, f"Plan de contrôle{suffixe} {cle_S} — publié le {lundi_S.strftime('%d/%m/%Y')}"),
            file_name=f"Plan_controle_{cle_S}{('_' + pp.PERIMETRES[perimetre]['fichier']) if pp.PERIMETRES[perimetre].get('fichier') else ''}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="dl_plan_controle")
        try:
            import plotly.graph_objects as go
            par_poste = m5_p.groupby(lambda e: str(info["poste_travail"].get(e, "?"))).sum().sort_values()
            par_poste = par_poste[par_poste >= 0.5].tail(20)
            fig = go.Figure(go.Bar(x=par_poste.values, y=par_poste.index, orientation="h",
                                   marker_color="#2a78d6", text=[f"{v:.1f}" for v in par_poste.values],
                                   textposition="outside"))
            fig.update_layout(title="Pannes probables par poste de travail (somme des probabilités)",
                              height=max(260, 28 * len(par_poste) + 80), margin=dict(l=10, r=30, t=50, b=10),
                              xaxis_title="Nombre de pannes attendu", plot_bgcolor="white")
            st.plotly_chart(fig, use_container_width=True)
        except Exception:
            pass

    # ══ 2. Résultats réels de la semaine précédente ═══════════════════════
    st.markdown("---")
    cle_prec = pp.semaine_label(lundi_S - pd.Timedelta(days=7))
    st.markdown(f"#### ✅ Résultat réel de la prédictive — {cle_prec}")
    ev = reg["evaluations"].get(cle_prec)
    if ev is None:
        if cle_prec not in reg["predictions"]:
            st.info(f"Aucun plan n'avait été publié pour {cle_prec} : le premier résultat réel "
                    f"s'affichera le {_fmt(lundi_S1)}, pour le plan de {cle_S}.")
        else:
            st.info(f"En attente d'une extraction couvrant toute la semaine {cle_prec} "
                    f"(jusqu'au dimanche {(lundi_S - pd.Timedelta(days=1)).strftime('%d/%m')}).")
    else:
        pr = reg["predictions"][cle_prec]
        reel = set(ev["pannes_reelles_liste"])
        sc = pd.Series(pr["scores"]["M5"])
        if postes:
            sc = sc[[str(info["poste_travail"].get(e, "")) in set(postes) for e in sc.index]]
            reel = {e for e in reel if str(info["poste_travail"].get(e, "")) in set(postes)}
        top = list(sc.sort_values(ascending=False).head(pp.TOP_N).index)
        hits = set(top) & reel
        r1, r2, r3, r4 = st.columns(4)
        r1.metric("Pannes réelles", f"{len(reel)}", "postes techniques (périmètre)", delta_color="off")
        r2.metric("Pannes prévues", f"{sc.sum():.0f}", "somme des probabilités", delta_color="off")
        r3.metric("Précision du plan", f"{len(hits) / max(len(top), 1) * 100:.0f} %",
                  f"{len(hits)} / {len(top)} postes du plan en panne", delta_color="off")
        r4.metric("Pannes détectées", f"{len(hits) / max(len(reel), 1) * 100:.0f} %",
                  f"{len(hits)} / {len(reel)} pannes étaient dans le plan", delta_color="off")
        with st.expander("Détail : pannes détectées et pannes non prévues"):
            det = pd.DataFrame({"Poste technique": sorted(hits)})
            det["Désignation"] = det["Poste technique"].map(lambda e: info["designation"].get(e, ""))
            manq = pd.DataFrame({"Poste technique": sorted(reel - set(top))})
            manq["Désignation"] = manq["Poste technique"].map(lambda e: info["designation"].get(e, ""))
            a, b = st.columns(2)
            a.markdown("**Détectées par le plan**"); a.dataframe(det, hide_index=True, use_container_width=True)
            b.markdown("**Non prévues**"); b.dataframe(manq, hide_index=True, use_container_width=True)

    # ══ 2 bis. Efficacité enregistrée chaque début de semaine + tendance ══
    _section_efficacite(reg, evol, cle_S, lundi_S)

    # ══ 3. Évolution des 5 modèles ════════════════════════════════════════
    st.markdown("---")
    st.markdown("#### 📈 Évolution des 5 modèles")
    if not evol:
        st.info("Pas encore de semaine évaluée.")
    else:
        crit = st.radio("Indicateur", ["Précision du plan", "Pannes détectées"], horizontal=True,
                        key="pred_crit")
        k = "precision" if crit == "Précision du plan" else "detection"
        try:
            import plotly.graph_objects as go
            fig = go.Figure()
            couleurs = {"M1": "#94a3b8", "M2": "#eb6834", "M3": "#1baf7a", "M4": "#eda100", "M5": "#2a78d6"}
            x = [e["semaine"].split("-")[1] + (" (réel)" if e.get("type") == "réel" else "") for e in evol]
            for m, nom in pp.MODELES.items():
                fig.add_trace(go.Scatter(
                    x=x, y=[e["modeles"][m][k] * 100 for e in evol], name=f"{m} · {nom}", mode="lines+markers",
                    line=dict(color=couleurs[m], width=4 if m == "M5" else 2)))
            n_retro = sum(1 for e in evol if e.get("type") != "réel")
            if 0 < n_retro < len(evol):
                fig.add_vline(x=n_retro - 0.5, line_dash="dash", line_color="#64748b")
            fig.update_layout(height=380, yaxis_title=f"{crit} (%)", plot_bgcolor="white",
                              margin=dict(l=10, r=10, t=30, b=10), legend=dict(orientation="h", y=-0.2))
            st.plotly_chart(fig, use_container_width=True)
        except Exception:
            pass
        lignes = []
        for m, nom in pp.MODELES.items():
            vals_p = [e["modeles"][m]["precision"] for e in evol[-8:]]
            vals_d = [e["modeles"][m]["detection"] for e in evol[-8:]]
            lignes.append({"Modèle": f"{m} · {nom}",
                           "Précision moyenne (8 sem.)": f"{sum(vals_p) / len(vals_p) * 100:.1f} %",
                           "Pannes détectées (8 sem.)": f"{sum(vals_d) / len(vals_d) * 100:.1f} %",
                           "Poids dans M5 cette semaine": (f"{pred['poids_M5'].get(m, 0):.2f}" if m != "M5" else "—")})
        st.dataframe(pd.DataFrame(lignes), hide_index=True, use_container_width=True)
        st.caption("Semaines « rétro-test » : rejouées sur l'historique comme si la prédiction avait été faite "
                   "le lundi. Semaines « réel » : plans publiés dans l'application puis comparés aux pannes "
                   "réelles. M5 donne plus de poids aux modèles qui ont le mieux prédit les dernières semaines.")

    # ══ 4. Méthode (masquée) ═════════════════════════════════════════════
    n_tr = f"{pred['entrainement'].get('n_train', 0):,}".replace(",", " ")
    n_pos = f"{pred['entrainement'].get('pos_train', 0):,}".replace(",", " ")
    with st.expander("⚙️ Méthode et calculs"):
        st.markdown(f"""
**Panne** : sur le **poste technique exact** saisi dans SAP (4 à 7 niveaux), un OT **ZCOR** de l'atelier
(avec ou sans avis) ou un avis **ZC** non rattaché à un de ces OT, décrivant un problème **mécanique, électrique ou d'instrumentation** : roulements,
alignement, accouplement, courroies / transmission, machines tournantes, moteur et alimentation,
capteurs et étalonnage, vannes, remise en état. **Exclus** : fuites et étanchéité, chaudronnerie
(soudure, tôle), tuyauterie, génie civil, bouchage et travaux annexes (échafaudage, peinture, nettoyage…).

**Variables** : pannes des 1, 2, 4, 8, 13, 26 et 52 dernières semaines, temps depuis la dernière panne,
intervalle moyen entre pannes (MTBF), pannes du poste technique parent, niveau du poste technique, préventif (OT ZPRV), correctif et inspections (avis ZO / ZI) récents, division.

**Modèles** (ré-entraînés chaque semaine sur tout l'historique) : M1 historique, M2 régression logistique,
M3 Random Forest, M4 Gradient Boosting, M5 ensemble adaptatif.

**Entraînement de cette semaine** : {n_tr} exemples
({n_pos} semaines-poste technique en panne) · plan généré le
{pred['genere_le']} · registre : {cache['source']}{(' · ' + cache['msg']) if cache['msg'] else ''}.
""")
